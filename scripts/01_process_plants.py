#!/usr/bin/env python3
"""Process the Philippine Power Plant Inventory XLSX into a canonical plants.geojson.

Unifies Master_Inventory (grid) and SPUG (off-grid) sheets into one schema:
  name, owner, parent, fuel, tech, status, island, region, province, municipality,
  lat, lon, capacity_mw, dependable_mw, source
"""
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data_sources" / "Philippine Power Plant Inventory for RE Compass.xlsx"
OUT = ROOT / "app" / "public" / "data" / "plants.geojson"

FUEL_MAP = {
    "high-speed diesel": "Oil-Based",
    "diesel": "Oil-Based",
    "solar": "Solar",
    "wind": "Wind",
    "hydro": "Hydro",
    "geothermal": "Geothermal",
    "biomass": "Biomass",
    "battery": "BESS",
    "bess": "BESS",
}


def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v != v:  # NaN
        return None
    s = str(v).strip()
    return s if s else None


def num(v):
    if v is None:
        return None
    try:
        f = float(str(v).replace(",", "").strip())
        return None if f != f else f
    except (ValueError, TypeError):
        return None


def guess_fuel(tech):
    if not tech:
        return None
    t = tech.lower()
    for k, v in FUEL_MAP.items():
        if k in t:
            return v
    return None


def master_rows():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Master_Inventory"]
    rows = list(ws.iter_rows(values_only=True))
    cols = [clean(c) for c in rows[0]]
    idx = {c: i for i, c in enumerate(cols) if c}
    for r in rows[1:]:
        if not clean(r[idx["Facility Name"]]):
            continue
        props = {c: clean(r[i]) for c, i in idx.items()}
        yield props


def spug_rows():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["SPUG"]
    rows = list(ws.iter_rows(values_only=True))
    cols = [clean(c) for c in rows[0]]
    idx = {c: i for i, c in enumerate(cols) if c}
    for r in rows[1:]:
        if not clean(r[idx["Facility Name"]]):
            continue
        props = {c: clean(r[i]) for c, i in idx.items()}
        yield props


def main():
    feats = []
    skipped = 0
    seen = set()

    for p in master_rows():
        name = p["Facility Name"]
        lat = num(p.get("Latitude (°N)"))
        lon = num(p.get("Longitude (°E)"))
        key = (name, p.get("Province"))
        rec = {
            "name": name,
            "owner": p.get("Operating SPV"),
            "parent": p.get("Parent Developer"),
            "fuel": p.get("Fuel Type"),
            "tech": p.get("Technology Sub-Type"),
            "status": p.get("Status"),
            "island": p.get("Grid"),
            "region": p.get("Region"),
            "province": p.get("Province"),
            "municipality": p.get("Municipality"),
            "capacity_mw": num(p.get("Installed (MW)")),
            "dependable_mw": num(p.get("Dependable (MW)")),
            "cod": p.get("Actual / Target COD"),
            "permitting": p.get("Permitting Framework"),
            "market": p.get("Contract / Market Type"),
            "source": "DOE grid inventory",
        }
        if lat is None or lon is None or not (4 < lat < 21.5) or not (115 < lon < 129.5):
            skipped += 1
            feats.append({"type": "Feature", "geometry": None, "properties": {**rec, "geocode_status": "missing_or_invalid_coords"}})
            continue
        if key in seen:
            continue
        seen.add(key)
        feats.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [round(lon, 6), round(lat, 6)]},
                "properties": {**rec, "geocode_status": "ok"},
            }
        )

    for p in spug_rows():
        name = p["Facility Name"]
        lat = num(p.get("Latitude (°N)"))
        lon = num(p.get("Longitude (°E)"))
        tech = p.get("Prime Mover & Technology") or p.get("Prime Mover & Techno")
        mw = num(p.get("Installed (MW)"))
        if mw is None:
            kw = num(p.get("Installed (kW)"))
            mw = round(kw / 1000, 3) if kw is not None else None
        dep = num(p.get("Dependable (MW)"))
        if dep is None:
            kw = num(p.get("Dependable (kW)"))
            dep = round(kw / 1000, 3) if kw is not None else None
        rec = {
            "name": name,
            "owner": p.get("SPUG Operational Division") or p.get("SPUG Operational Div"),
            "parent": None,
            "fuel": guess_fuel(tech) or "Oil-Based",
            "tech": tech,
            "status": "Operational",
            "island": p.get("Island Group"),
            "region": None,
            "province": p.get("Province"),
            "municipality": p.get("Municipality"),
            "barangay": p.get("Host Barangay"),
            "capacity_mw": mw,
            "dependable_mw": dep,
            "regulatory_status": p.get("Regulatory Status"),
            "source": "DOE SPUG (off-grid)",
        }
        if lat is None or lon is None or not (4 < lat < 21.5) or not (115 < lon < 129.5):
            skipped += 1
            feats.append({"type": "Feature", "geometry": None, "properties": {**rec, "geocode_status": "missing_or_invalid_coords"}})
            continue
        key = (name, p.get("Province"))
        if key in seen:
            continue
        seen.add(key)
        feats.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [round(lon, 6), round(lat, 6)]},
                "properties": {**rec, "geocode_status": "ok"},
            }
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"type": "FeatureCollection", "features": feats}, ensure_ascii=False))
    mapped = sum(1 for f in feats if f["geometry"])
    print(f"Wrote {len(feats)} features -> {OUT}")
    print(f"  with geometry: {mapped}, without: {len(feats) - mapped}")

    from collections import Counter

    withg = [f["properties"] for f in feats if f["geometry"]]
    print("fuel:", Counter(p["fuel"] or "?" for p in withg).most_common())
    print("status:", Counter(p["status"] or "?" for p in withg).most_common())
    print("island:", Counter(p["island"] or "?" for p in withg).most_common())


if __name__ == "__main__":
    main()
